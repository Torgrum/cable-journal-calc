"""Вспомогательные инструменты для работы с Excel-шаблонами."""
import re
import io
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.cell.cell import MergedCell


# ==================== Парсинг диапазонов ====================

def parse_range(range_str):
    s = range_str.strip().upper().replace(" ", "")
    if not s:
        raise ValueError("Диапазон пустой")
    if ":" not in s:
        s = f"{s}:{s}"
    m = re.match(r"^([A-Z]+)(\d+):([A-Z]+)(\d+)$", s)
    if not m:
        raise ValueError("Формат: 'B5:J6' (например)")
    c1 = column_index_from_string(m.group(1))
    r1 = int(m.group(2))
    c2 = column_index_from_string(m.group(3))
    r2 = int(m.group(4))
    return min(r1, r2), min(c1, c2), max(r1, r2), max(c1, c2)


# ==================== Работа с merged cells ====================

def get_merged_value(ws, row, col):
    cell = ws.cell(row=row, column=col)
    if isinstance(cell, MergedCell):
        for rng in ws.merged_cells.ranges:
            if cell.coordinate in rng:
                return ws.cell(row=rng.min_row, column=rng.min_col).value
    return cell.value


def get_merged_ranges_in_range(ws, r1, c1, r2, c2):
    result = []
    for rng in ws.merged_cells.ranges:
        if (rng.min_row <= r2 and rng.max_row >= r1 and
                rng.min_col <= c2 and rng.max_col >= c1):
            local_r1, local_c1 = max(rng.min_row, r1), max(rng.min_col, c1)
            local_r2, local_c2 = min(rng.max_row, r2), min(rng.max_col, c2)
            value = ws.cell(row=rng.min_row, column=rng.min_col).value
            result.append({
                "r1": local_r1, "c1": local_c1,
                "r2": local_r2, "c2": local_c2,
                "rowspan": local_r2 - local_r1 + 1,
                "colspan": local_c2 - local_c1 + 1,
                "value": value,
            })
    return result


# ==================== HTML-превью заголовка (СИНХРОНИЗИРОВАННЫЙ) ====================

# Палитра цветов для разных объединённых групп
GROUP_COLORS = [
    "#dbeafe",  # голубой
    "#fef3c7",  # жёлтый
    "#dcfce7",  # зелёный
    "#fce7f3",  # розовый
    "#e0e7ff",  # индиго
    "#fed7aa",  # оранжевый
    "#ddd6fe",  # фиолетовый
    "#ccfbf1",  # бирюзовый
]


def build_header_html(r1, c1, r2, c2, merged_ranges, col_widths_px):
    """
    Строит HTML-заголовок, синхронизированный по ширине с data_editor.

    col_widths_px: список ширин в пикселях для каждой колонки (c1..c2).
    """
    nrows = r2 - r1 + 1
    ncols = c2 - c1 + 1
    covered = [[False] * ncols for _ in range(nrows)]

    # Присваиваем цвета группам объединений
    color_map = {}
    color_idx = 0
    for mr in merged_ranges:
        if mr["colspan"] > 1 or mr["rowspan"] > 1:
            key = (mr["r1"], mr["c1"])
            if key not in color_map:
                color_map[key] = GROUP_COLORS[color_idx % len(GROUP_COLORS)]
                color_idx += 1

    html = (
        '<table style="border-collapse: collapse; margin: 0 0 4px 0; '
        'font-family: system-ui, sans-serif; font-size: 13px; '
        'table-layout: fixed; width: max-content;">'
    )

    # Определяем ширины колонок через <col>
    html += "<colgroup>"
    for w in col_widths_px:
        html += f'<col style="width: {w}px;">'
    html += "</colgroup>"

    for ri in range(nrows):
        html += "<tr>"
        for ci in range(ncols):
            if covered[ri][ci]:
                continue
            abs_r, abs_c = r1 + ri, c1 + ci
            found = next(
                (mr for mr in merged_ranges if mr["r1"] == abs_r and mr["c1"] == abs_c),
                None,
            )

            if found:
                # Отмечаем покрытые ячейки
                for dr in range(found["rowspan"]):
                    for dc in range(found["colspan"]):
                        if ri + dr < nrows and ci + dc < ncols:
                            covered[ri + dr][ci + dc] = True

                val = str(found["value"]) if found["value"] not in (None, "") else ""
                bg = color_map.get((found["r1"], found["c1"]), "#e7f3ff")
                is_merged = found["colspan"] > 1 or found["rowspan"] > 1
                border_left = "3px solid #1e40af" if is_merged and ci == 0 else "1px solid #888"

                html += (
                    f'<td rowspan="{found["rowspan"]}" colspan="{found["colspan"]}" '
                    f'style="border: 1px solid #888; border-left: {border_left}; '
                    f'padding: 8px 6px; background: {bg}; font-weight: 600; '
                    f'text-align: center; vertical-align: middle; '
                    f'color: #1f2937; white-space: normal; word-break: break-word;">{val}</td>'
                )
            else:
                # Обычная ячейка
                val_raw = ""
                html += (
                    f'<td style="border: 1px solid #888; padding: 8px 6px; '
                    f'text-align: center; vertical-align: middle; '
                    f'background: #f9fafb; color: #6b7280; font-weight: 500;">{val_raw}</td>'
                )
        html += "</tr>"
    html += "</table>"
    return html


# ==================== Парсинг шаблона ====================

def parse_template(file_bytes, sheet_name, header_range_str):
    wb = load_workbook(io.BytesIO(file_bytes))
    ws = wb[sheet_name]
    r1, c1, r2, c2 = parse_range(header_range_str)

    merged_ranges = get_merged_ranges_in_range(ws, r1, c1, r2, c2)

    def top_left_for(row, col):
        for mr in merged_ranges:
            if mr["r1"] <= row <= mr["r2"] and mr["c1"] <= col <= mr["c2"]:
                return mr["r1"], mr["c1"]
        return row, col

    columns = {}
    for col in range(c1, c2 + 1):
        parts_by_group = {}
        for r in range(r1, r2 + 1):
            tl = top_left_for(r, col)
            val = get_merged_value(ws, r, col)
            if val not in (None, ""):
                parts_by_group.setdefault(tl, []).append(str(val).strip().replace("\n", " "))

        seen = []
        for tl in sorted(parts_by_group.keys()):
            if parts_by_group[tl]:
                v = parts_by_group[tl][0]
                if v not in seen:
                    seen.append(v)

        if seen:
            name = " / ".join(seen)
            final_name, counter = name, 1
            while final_name in columns:
                final_name = f"{name} [{get_column_letter(col)}]"
                counter += 1
                if counter > 10:
                    break
            columns[final_name] = col

    if not columns:
        raise ValueError(f"В диапазоне '{header_range_str}' нет заголовков")

    data_start_row = r2 + 1
    types = {}
    for name, idx in columns.items():
        sample = get_merged_value(ws, data_start_row, idx)
        types[name] = "number" if isinstance(sample, (int, float)) and not isinstance(sample, bool) else "text"

    return {
        "ws_name": ws.title,
        "header_range": f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}",
        "header_coords": (r1, c1, r2, c2),
        "data_start_row": data_start_row,
        "columns": columns,
        "types": types,
        "merged_ranges": merged_ranges,
    }


# ==================== Экспорт в Excel ====================

def export_to_excel(template_bytes, ws_name, data_start_row, columns_map, df):
    output = io.BytesIO()
    wb = load_workbook(io.BytesIO(template_bytes))
    ws = wb[ws_name]

    for r in range(data_start_row, data_start_row + 2000):
        all_empty = True
        for c in columns_map.values():
            cell = ws.cell(row=r, column=c)
            v = cell.value
            if v not in (None, "") and not (isinstance(v, str) and v.startswith("=")):
                all_empty = False
        if all_empty:
            break
        for c in columns_map.values():
            cell = ws.cell(row=r, column=c)
            if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                cell.value = None

    for i, (_, row) in enumerate(df.iterrows()):
        for col_name, val in row.items():
            if col_name in columns_map:
                cell = ws.cell(row=data_start_row + i, column=columns_map[col_name])
                if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                    try:
                        cell.value = None if _is_empty(val) else _to_python(val)
                    except Exception:
                        cell.value = None

    wb.save(output)
    return output.getvalue()


def _is_empty(val):
    if val is None or val == "":
        return True
    try:
        import pandas as pd
        return pd.isna(val)
    except Exception:
        return False


def _to_python(val):
    import numpy as np
    if isinstance(val, (np.integer,)):
        return int(val)
    if isinstance(val, (np.floating,)):
        return float(val)
    if isinstance(val, (np.bool_,)):
        return bool(val)
    return val
