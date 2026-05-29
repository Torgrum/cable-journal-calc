"""Упрощённый калькулятор кабельного журнала (без AgGrid)."""
import streamlit as st
import pandas as pd
import io
import json
from openpyxl.utils import get_column_letter
from openpyxl import load_workbook

from excel_tools import parse_template, export_to_excel

st.set_page_config(page_title="Универсальный КЖ", page_icon="📋", layout="wide")
st.title("📋 Универсальный калькулятор кабельного журнала")

if "template_bytes" not in st.session_state:
    st.session_state.template_bytes = None
if "config" not in st.session_state:
    st.session_state.config = None
if "df_initial" not in st.session_state:
    st.session_state.df_initial = None

def reset_editor():
    if "data_editor" in st.session_state:
        del st.session_state["data_editor"]

with st.sidebar:
    st.header("⚙️ Настройки")
    uploaded = st.file_uploader("📄 Загрузить шаблон .xlsx", type=["xlsx"])
    if uploaded is not None:
        new_bytes = uploaded.getvalue()
        if st.session_state.template_bytes != new_bytes:
            st.session_state.template_bytes = new_bytes
            st.session_state.config = None
            st.session_state.df_initial = None

    if st.session_state.template_bytes is not None:
        wb_tmp = load_workbook(io.BytesIO(st.session_state.template_bytes), read_only=True)
        selected_sheet = st.selectbox("Лист", wb_tmp.sheetnames)
        wb_tmp.close()

        st.markdown("**Диапазон заголовка** (напр. `B5:J6`)")
        header_range = st.text_input("Диапазон", value="A1:E1", placeholder="B5:J6")

        if st.button("🔄 Применить шаблон", type="primary"):
            try:
                cfg = parse_template(st.session_state.template_bytes, selected_sheet, header_range)
                st.session_state.config = cfg
                st.session_state.df_initial = pd.DataFrame(columns=list(cfg["columns"].keys()))
                reset_editor()
                st.success(f"✅ Шаблон распознан. Столбцов: {len(cfg['columns'])}")
                st.rerun()
            except Exception as e:
                st.error(f"Ошибка: {e}")

if st.session_state.config is None:
    st.info("👈 Загрузите шаблон `.xlsx` слева и укажите диапазон заголовков.")
    st.stop()

cfg = st.session_state.config
cols = list(cfg["columns"].keys())

if st.session_state.df_initial is None or list(st.session_state.df_initial.columns) != cols:
    st.session_state.df_initial = pd.DataFrame(columns=cols)

with st.expander("📐 Структура шаблона", expanded=False):
    st.write(f"**Лист:** `{cfg['ws_name']}` | **Диапазон:** `{cfg['header_range']}` | **Старт данных:** строка {cfg['data_start_row']}")
    if cfg["merged_ranges"]:
        st.info(f"🔗 Объединений в заголовке: **{len(cfg['merged_ranges'])}** (сохранятся при экспорте)")

st.subheader("📝 Заполнение данных")
st.caption("💡 Данные сохраняются автоматически. Перед закрытием используйте **💾 Сохранить проект**.")

col_add, col_clear, _ = st.columns([1, 1, 4])
with col_add:
    if st.button("➕ Добавить строку"):
        new_row = pd.DataFrame([{c: None for c in cols}])
        st.session_state.df_initial = pd.concat([st.session_state.df_initial, new_row], ignore_index=True)
        reset_editor()
        st.rerun()
with col_clear:
    if st.button("🗑️ Очистить всё"):
        st.session_state.df_initial = pd.DataFrame(columns=cols)
        reset_editor()
        st.rerun()

col_cfg = {}
for n in cols:
    if cfg["types"][n] == "number":
        col_cfg[n] = st.column_config.NumberColumn(n, step=0.5, width="medium")
    else:
        col_cfg[n] = st.column_config.TextColumn(n, width="medium")

edited_df = st.data_editor(
    st.session_state.df_initial,
    column_config=col_cfg,
    column_order=cols,
    hide_index=True,
    use_container_width=True,
    num_rows="dynamic",
    key="data_editor",
)

st.session_state.df_initial = edited_df
df_clean = edited_df.copy().replace('', pd.NA).dropna(how='all').reset_index(drop=True)

if not df_clean.empty:
    st.subheader("📊 Сводка")
    c1, c2 = st.columns(2)
    c1.metric("Строк заполнено", len(df_clean))
    num_cols = df_clean.select_dtypes(include=['number']).columns
    if len(num_cols) > 0:
        c2.metric("Сумма по числам", f"{df_clean[num_cols].sum().sum():.2f}")

    st.divider()

    if st.button("📥 Сформировать Excel", type="primary"):
        try:
            result_bytes = export_to_excel(
                template_bytes=st.session_state.template_bytes,
                ws_name=cfg["ws_name"],
                data_start_row=cfg["data_start_row"],
                columns_map=cfg["columns"],
                df=df_clean,
            )
            st.download_button("⬇️ Скачать файл", result_bytes, "journal_filled.xlsx",
                               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                               use_container_width=True)
            st.success("✅ Файл готов! Объединения и стили шаблона сохранены.")
        except Exception as e:
            st.error(f"Ошибка экспорта: {e}")

    try:
        json_str = df_clean.to_json(orient="records", force_ascii=False, default_handler=str)
        st.download_button("💾 Сохранить проект (.json)", json_str.encode("utf-8"), "project.json", "application/json")
    except Exception:
        pass

    uploaded_project = st.file_uploader("📂 Загрузить сохранённый проект", type=["json"], key="proj_loader")
    if uploaded_project is not None:
        try:
            data = json.loads(uploaded_project.read().decode("utf-8"))
            loaded_df = pd.DataFrame(data)
            for c in cols:
                if c not in loaded_df.columns: loaded_df[c] = None
            st.session_state.df_initial = loaded_df[cols]
            reset_editor()
            st.success(f"✅ Загружено {len(loaded_df)} строк")
            st.rerun()
        except Exception as e:
            st.error(f"Ошибка: {e}")
else:
    st.info("Таблица пуста. Нажмите **➕ Добавить строку** выше.")
